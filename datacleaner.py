#!/usr/bin/env python3
"""
datacleaner.py - clean messy CSV/Excel files, remove duplicates, produce a report.

Examples
    python datacleaner.py customers.xlsx
    python datacleaner.py customers.csv --keys email --outdir out/
    python datacleaner.py sales.xlsx --sheet Q3 --keys "email,first_name+last_name" --phone-prefix +92

Outputs (in --outdir, default ./output)
    <name>_clean.xlsx    Clean Data | Removed Duplicates | Issues | Cleaning Summary
    <name>_clean.csv     the clean rows, ready to import elsewhere
    <name>_report.md     the same summary as a plain-text report

The original file is never modified.
"""
from __future__ import annotations

import argparse
import csv
import io
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from dateutil import parser as dateparser
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

__version__ = "1.0"

FONT = "Arial"
NULL_TOKENS = {"", "n/a", "na", "n.a.", "null", "none", "nan", "nil", "-", "--", "?", "unknown", "tbd"}
EMAIL_RE = re.compile(r"^[a-z0-9._%+\-]+@[a-z0-9\-]+(\.[a-z0-9\-]+)*\.[a-z]{2,}$")
FOOTER_RE = re.compile(r"^(grand\s+)?(total|subtotal|sum|generated|printed|end of|records?\b)", re.I)
ISO_RE = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})([ T].*)?$")
DATE_LIKE_RE = re.compile(r"^\d{1,4}[-/. ]\d{1,2}[-/. ]\d{1,4}$|[A-Za-z]{3,}")
ID_LIKE_RE = re.compile(r"(^|_)(id|code|zip|postal|sku|ref|reference|cnic|nic|passport|invoice|order_no|no|number)($|_)")
COUNTRY_MAP = {
    "pk": "Pakistan", "pak": "Pakistan", "pakistan": "Pakistan",
    "us": "United States", "usa": "United States", "u.s.a": "United States", "u.s": "United States",
    "united states": "United States", "united states of america": "United States",
    "uk": "United Kingdom", "u.k": "United Kingdom", "gb": "United Kingdom", "england": "United Kingdom",
    "united kingdom": "United Kingdom", "uae": "United Arab Emirates", "united arab emirates": "United Arab Emirates",
    "ksa": "Saudi Arabia", "saudi arabia": "Saudi Arabia", "ca": "Canada", "canada": "Canada",
}


# --------------------------------------------------------------------------- loading
def read_raw(path: Path, sheet=None) -> pd.DataFrame:
    """Read the file as a grid of strings with no header assumptions."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xls"):
        if sheet is not None and str(sheet).isdigit():
            sheet = int(sheet)
        return pd.read_excel(path, sheet_name=sheet if sheet is not None else 0,
                             header=None, dtype=str, keep_default_na=False)
    if suffix in (".csv", ".tsv", ".txt"):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        width = max((len(r) for r in rows), default=0)
        return pd.DataFrame([r + [""] * (width - len(r)) for r in rows], dtype=str)
    raise SystemExit(f"Unsupported file type: {suffix} (use .xlsx, .xls, .csv, .tsv)")


def clean_cell(v):
    """Trim, collapse whitespace, and turn 'N/A'-style tokens into None."""
    if v is None or (isinstance(v, float) and v != v):
        return None
    s = str(v).replace("\u00a0", " ").replace("\u200b", "").replace("\ufeff", "")
    s = re.sub(r"\s+", " ", s).strip()
    return None if s.lower() in NULL_TOKENS else s


def snake(s: str) -> str:
    s = re.sub(r"[^0-9a-zA-Z]+", "_", str(s).strip().lower()).strip("_")
    s = s.replace("e_mail", "email")
    return s or "column"


def detect_header(raw: pd.DataFrame) -> int:
    """First row (within the top 25) that is nearly as full as the fullest row."""
    counts = [sum(clean_cell(v) is not None for v in row) for row in raw.head(25).itertuples(index=False)]
    if not counts or max(counts) == 0:
        raise SystemExit("The file looks empty.")
    threshold = max(2, 0.8 * max(counts))
    return next(i for i, c in enumerate(counts) if c >= threshold)


# --------------------------------------------------------------------------- parsers
def parse_date(s: str, dayfirst: bool):
    if not DATE_LIKE_RE.search(s):
        return None
    m = ISO_RE.match(s)
    try:
        if m:
            return date(int(m[1]), int(m[2]), int(m[3]))
        return dateparser.parse(s, dayfirst=dayfirst, default=datetime(1900, 1, 1)).date()
    except (ValueError, OverflowError):
        return None


def parse_number(s: str):
    t = s.strip()
    neg = t.startswith("(") and t.endswith(")")
    if neg:
        t = t[1:-1]
    t = re.sub(r"(?i)(pkr|usd|inr|eur|gbp|rs\.?)|[$£€₨]", "", t)
    t = t.replace(",", "").replace(" ", "")
    if not re.fullmatch(r"-?\d+(\.\d+)?", t):
        return None
    v = float(t)
    v = -v if neg else v
    return int(v) if v.is_integer() else round(v, 4)


def fmt_phone(s: str, prefix: str):
    d = re.sub(r"[^\d+]", "", s)
    if d.startswith("+"):
        d = "+" + re.sub(r"\D", "", d)
    else:
        d = re.sub(r"\D", "", d)
        if d.startswith("00"):
            d = "+" + d[2:]
        elif prefix and d.startswith("0"):
            d = prefix + d[1:]
        elif prefix and len(d) == 10:
            d = prefix + d
        elif prefix and d.startswith(prefix.lstrip("+")):
            d = "+" + d
    if not 8 <= len(re.sub(r"\D", "", d)) <= 15:
        return s, "Phone number has an unusual length"
    return d, None


def fmt_email(s: str):
    e = s.lower().replace("mailto:", "").strip(" .,;<>")
    return (e, None) if EMAIL_RE.match(e) else (e, "Invalid email format")


def smart_title(s: str) -> str:
    """Title-case only text that is ALL CAPS or all lower; leave McDonald / O'Brien alone."""
    if s.islower() or (s.isupper() and (" " in s or len(s) > 5)):
        return re.sub(r"[A-Za-z]+('[A-Za-z]+)?", lambda m: m.group(0).capitalize(), s)
    return s


def as_text(v) -> str:
    if v is None or (not isinstance(v, (list, dict)) and pd.isna(v)):
        return ""
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


# --------------------------------------------------------------------------- column typing
def infer_kind(name: str, values: list[str], dayfirst: bool) -> str:
    if not values:
        return "empty"
    if "email" in name:
        return "email"
    if re.search(r"phone|mobile|tel|cell|whatsapp|fax", name):
        return "phone"
    if ID_LIKE_RE.search(name):
        return "text"
    n = len(values)
    if re.search(r"date|dob|birth|created|updated|joined|signup", name) or \
            sum(parse_date(v, dayfirst) is not None for v in values) / n >= 0.8:
        return "date"
    leading_zero = any(len(v) > 1 and v[0] == "0" and "." not in v for v in values)
    if not leading_zero and (
            re.search(r"amount|price|total|revenue|cost|salary|balance|qty|quantity|score|age", name)
            or sum(parse_number(v) is not None for v in values) / n >= 0.8):
        return "number"
    if "country" in name:
        return "country"
    if re.search(r"name|city|town|state|province|company|organi[sz]ation|address|street", name):
        return "title"
    if n >= 8 and len({v.lower() for v in values}) <= 12 and len({v.lower() for v in values}) <= n / 2:
        return "category"  # e.g. Active / active / ACTIVE -> one spelling
    return "text"


def format_value(v: str, kind: str, opts):
    """Return (clean_value, issue_or_None)."""
    if kind == "email":
        return fmt_email(v)
    if kind == "phone":
        return fmt_phone(v, opts.phone_prefix)
    if kind == "date":
        d = parse_date(v, opts.dayfirst)
        return (d, None) if d else (v, "Could not read as a date")
    if kind == "number":
        n = parse_number(v)
        return (n, None) if n is not None else (v, "Could not read as a number")
    if kind == "country":
        return COUNTRY_MAP.get(v.lower().replace(".", "").strip(), smart_title(v)), None
    if kind == "title":
        return smart_title(v), None
    return v, None


# --------------------------------------------------------------------------- dedupe
def key_part(v, kind):
    if v is None or pd.isna(v):
        return None
    s = as_text(v).lower().strip()
    if kind == "phone":
        s = re.sub(r"\D", "", s)[-10:]
    return s or None


def dedupe(df: pd.DataFrame, kinds: dict, key_specs: list[list[str]]):
    data_cols = [c for c in df.columns if c != "source_row"]
    removed, filled = [], 0

    # 1) exact duplicates (ignoring case)
    seen, drop = {}, []
    for idx, row in df.iterrows():
        sig = tuple(as_text(row[c]).lower() for c in data_cols)
        if sig in seen:
            removed.append({**row.to_dict(), "removed_reason": "Exact duplicate",
                            "kept_source_row": df.at[seen[sig], "source_row"]})
            drop.append(idx)
        else:
            seen[sig] = idx
    df = df.drop(index=drop)

    # 2) key duplicates: keep the most complete row and fill its gaps from the others
    for spec in key_specs:
        groups = defaultdict(list)
        for idx, row in df.iterrows():
            parts = [key_part(row[c], kinds.get(c)) for c in spec]
            if all(p is not None for p in parts):
                groups[tuple(parts)].append(idx)
        drop = []
        for idxs in groups.values():
            if len(idxs) < 2:
                continue
            best = max(idxs, key=lambda i: (df.loc[i, data_cols].notna().sum(), -df.at[i, "source_row"]))
            for i in idxs:
                if i == best:
                    continue
                for c in data_cols:
                    if pd.isna(df.at[best, c]) and not pd.isna(df.at[i, c]):
                        df.at[best, c] = df.at[i, c]
                        filled += 1
                removed.append({**df.loc[i].to_dict(), "removed_reason": f"Duplicate {' + '.join(spec)}",
                                "kept_source_row": df.at[best, "source_row"]})
                drop.append(i)
        df = df.drop(index=drop)
    return df, removed, filled


# --------------------------------------------------------------------------- pipeline
def clean_file(path: Path, opts) -> dict:
    raw = read_raw(path, opts.sheet)
    hdr = detect_header(raw)

    names, seen = [], defaultdict(int)
    for j, v in enumerate(raw.iloc[hdr]):
        base = snake(clean_cell(v)) if clean_cell(v) else f"column_{j + 1}"
        seen[base] += 1
        names.append(base if seen[base] == 1 else f"{base}_{seen[base]}")

    body = raw.iloc[hdr + 1:].copy()
    body.columns = names
    body.insert(0, "source_row", [i + 1 for i in body.index])  # row number as seen in Excel
    orig = body.copy()
    cells = body[names].astype(object).map(clean_cell)
    cells = cells.astype(object).where(cells.notna(), None)  # pandas 3 turns None into NaN

    # drop blank rows and footer rows ("Total records: 40")
    n_source = len(cells)
    keep_mask = []
    for _, row in cells.iterrows():
        filled = [v for v in row if not pd.isna(v)]
        keep_mask.append(bool(filled) and not (len(filled) <= 2 and FOOTER_RE.match(filled[0])))
    keep = pd.Series(keep_mask, index=cells.index)
    n_blank = int((~keep).sum())
    cells, orig = cells[keep], orig[keep]

    empty_cols = [c for c in names if cells[c].isna().all()]
    cols = [c for c in names if c not in empty_cols]

    kinds, out, issues, col_stats = {}, {}, [], {}
    for c in cols:
        kinds[c] = infer_kind(c, [v for v in cells[c] if not pd.isna(v)], opts.dayfirst)
        vals, changed, examples, tokens = [], 0, [], 0
        canon = {}
        if kinds[c] == "category":
            counts = Counter(v for v in cells[c] if not pd.isna(v))
            for spelling, _ in counts.most_common():
                canon.setdefault(spelling.lower(), spelling)
        for idx in cells.index:
            o, v = orig.at[idx, c], cells.at[idx, c]
            if pd.isna(v):
                vals.append(None)
                if str(o).strip():
                    tokens += 1
                continue
            new, issue = (canon[v.lower()], None) if canon else format_value(v, kinds[c], opts)
            vals.append(new)
            if issue:
                issues.append({"source_row": orig.at[idx, "source_row"], "column": c, "value": v, "issue": issue})
            if as_text(new) != str(o):
                changed += 1
                if len(examples) < 2:
                    examples.append(f"{o!r} -> {as_text(new)!r}")
        out[c] = vals
        col_stats[c] = {"kind": kinds[c], "changed": changed, "tokens": tokens, "example": "; ".join(examples)}

    df = pd.DataFrame(out, index=cells.index, dtype=object)
    df.insert(0, "source_row", orig["source_row"].values)
    df = df.astype(object).where(df.notna(), None)

    # dedupe rules
    if opts.keys:
        key_specs = [[snake(p) for p in k.split("+")] for k in opts.keys.split(",") if k.strip()]
        bad = [c for s in key_specs for c in s if c not in cols]
        if bad:
            raise SystemExit(f"--keys refers to unknown column(s): {bad}. Available: {cols}")
    else:
        email_cols = [c for c in cols if kinds[c] == "email"]
        key_specs = [[email_cols[0]]] if email_cols else []

    df, removed, filled = dedupe(df, kinds, key_specs)
    kept_rows = set(df["source_row"])
    issues = [i for i in issues if i["source_row"] in kept_rows]

    for c in cols:
        col_stats[c]["missing"] = int(df[c].isna().sum())

    return {
        "path": path, "sheet": opts.sheet if opts.sheet is not None else "first sheet",
        "header_row": hdr + 1, "n_source": n_source, "n_blank": n_blank,
        "empty_cols": empty_cols, "cols": cols, "kinds": kinds, "col_stats": col_stats,
        "clean": df.reset_index(drop=True), "removed": removed, "issues": issues,
        "filled": filled, "key_specs": key_specs,
        "n_exact": sum(r["removed_reason"] == "Exact duplicate" for r in removed),
    }


# --------------------------------------------------------------------------- output
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")


def _style_table(ws, n_cols, widths_from=None):
    for j in range(1, n_cols + 1):
        c = ws.cell(row=1, column=j)
        c.font, c.fill = Font(name=FONT, bold=True, color="FFFFFF", size=10), HEAD_FILL
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j in range(1, n_cols + 1):
        longest = max((len(as_text(ws.cell(row=r, column=j).value)) for r in range(1, min(ws.max_row, 200) + 1)), default=8)
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 2, 10), 42)


def _write_rows(ws, header, rows, kinds=None):
    ws.append(header)
    for r in rows:
        ws.append([None if (v is not None and not isinstance(v, (list, dict)) and pd.isna(v)) else v for v in r])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            if isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
    if kinds:
        for j, h in enumerate(header, start=1):
            if kinds.get(h) == "number":
                fmt = "#,##0.00" if any(isinstance(ws.cell(row=r, column=j).value, float)
                                        for r in range(2, ws.max_row + 1)) else "#,##0"
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=j).number_format = fmt
    _style_table(ws, len(header))


def write_xlsx(res: dict, out_path: Path):
    wb = Workbook()
    df, cols = res["clean"], res["cols"]

    ws = wb.active
    ws.title = "Clean Data"
    _write_rows(ws, ["source_row"] + cols, df[["source_row"] + cols].itertuples(index=False), res["kinds"])
    row_of = {sr: i + 2 for i, sr in enumerate(df["source_row"])}
    col_of = {c: j + 2 for j, c in enumerate(cols)}
    for iss in res["issues"]:
        ws.cell(row=row_of[iss["source_row"]], column=col_of[iss["column"]]).fill = FLAG_FILL

    ws = wb.create_sheet("Removed Duplicates")
    rows = [[r["source_row"], r["removed_reason"], r["kept_source_row"]] + [r.get(c) for c in cols] for r in res["removed"]]
    _write_rows(ws, ["source_row", "removed_reason", "kept_source_row"] + cols, rows, res["kinds"])

    ws = wb.create_sheet("Issues")
    _write_rows(ws, ["source_row", "column", "value", "issue"],
                [[i["source_row"], i["column"], i["value"], i["issue"]] for i in res["issues"]])

    ws = wb.create_sheet("Cleaning Summary")
    f = lambda **k: Font(name=FONT, size=10, **k)
    ws["A1"], ws["A1"].font = "Data Cleaning Report", Font(name=FONT, size=14, bold=True, color="1F3864")
    ws["A2"], ws["A2"].font = f"Generated by datacleaner v{__version__}", f(italic=True, color="666666")

    def section(r, title):
        for j in range(1, 7):
            ws.cell(row=r, column=j).fill = SECTION_FILL
        ws.cell(row=r, column=1, value=title).font = f(bold=True)

    def line(r, label, value, note=None):
        ws.cell(row=r, column=1, value=label).font = f()
        ws.cell(row=r, column=2, value=value).font = f()
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="left")
        if note:
            ws.cell(row=r, column=3, value=note).font = f(italic=True, color="666666")

    section(4, "Source")
    line(5, "File", res["path"].name)
    line(6, "Sheet", str(res["sheet"]))
    line(7, "Header found at source row", res["header_row"], "Title lines above the header are ignored")
    line(8, "Processed", datetime.now().strftime("%Y-%m-%d %H:%M"))

    section(10, "Row accounting")
    line(11, "Data rows in source", res["n_source"])
    line(12, "Blank / footer rows removed", res["n_blank"])
    line(13, "Exact duplicates removed", "=COUNTIF('Removed Duplicates'!B:B,\"Exact duplicate\")")
    line(14, "Key-based duplicates removed", "=COUNTA('Removed Duplicates'!A:A)-1-B13")
    line(15, "Clean rows delivered", "=COUNTA('Clean Data'!A:A)-1")
    line(16, "Reconciliation (must be 0)", "=B11-B12-B13-B14-B15", "Every source row is accounted for")
    for r in (15, 16):
        ws.cell(row=r, column=1).font = ws.cell(row=r, column=2).font = f(bold=True)

    section(18, "Quality")
    line(19, "Cells flagged for review", "=COUNTA(Issues!A:A)-1", "Highlighted yellow in Clean Data; listed in Issues")
    line(20, "Gaps filled from duplicate records", res["filled"], "Missing values completed from the removed copy")
    line(21, "Empty columns dropped", ", ".join(res["empty_cols"]) or "none")
    rules = ["exact row match"] + [" + ".join(k) for k in res["key_specs"]]
    line(22, "Duplicate rules applied", "; ".join(rules), "Most complete record is kept")

    section(24, "Column report")
    for j, h in enumerate(["Column", "Detected type", "Cells changed", "Missing", "Completeness", "Example change"], 1):
        c = ws.cell(row=25, column=j, value=h)
        c.font, c.fill = Font(name=FONT, size=10, bold=True, color="FFFFFF"), HEAD_FILL
    for i, c in enumerate(cols):
        r, s = 26 + i, res["col_stats"][c]
        for j, v in enumerate([c, s["kind"], s["changed"], s["missing"], f"=IFERROR(1-D{r}/$B$15,0)", s["example"]], 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = f()
            if j == 5:
                cell.number_format = "0%"
    last = 26 + len(cols) + 1
    ws.cell(row=last, column=1, value="Notes").font = f(bold=True)
    notes = ["Cells changed counts all source rows before duplicate removal.",
             "Currency symbols/labels are stripped from numbers; assumes one currency per column.",
             "Dates are written as yyyy-mm-dd; day/month order is set with --monthfirst."]
    for k, n in enumerate(notes):
        ws.cell(row=last + 1 + k, column=1, value=n).font = f(italic=True, color="666666")
    for col, w in zip("ABCDEF", (34, 24, 16, 12, 14, 60)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False

    wb.save(out_path)


def write_markdown(res: dict, out_path: Path):
    n_keep, n_dup = len(res["clean"]), len(res["removed"])
    L = [f"# Data Cleaning Report - {res['path'].name}", "",
         f"_Generated {datetime.now():%Y-%m-%d %H:%M} by datacleaner v{__version__}_", "",
         "## Summary", "",
         f"- Data rows in source: **{res['n_source']}**",
         f"- Blank / footer rows removed: **{res['n_blank']}**",
         f"- Duplicates removed: **{n_dup}** ({res['n_exact']} exact, {n_dup - res['n_exact']} by key)",
         f"- Gaps filled from duplicate records: **{res['filled']}**",
         f"- Clean rows delivered: **{n_keep}**",
         f"- Cells flagged for review: **{len(res['issues'])}**", "",
         "## Columns", "", "| Column | Type | Cells changed | Missing | Example |", "|---|---|---|---|---|"]
    for c in res["cols"]:
        s = res["col_stats"][c]
        L.append(f"| {c} | {s['kind']} | {s['changed']} | {s['missing']} | {s['example'] or '-'} |")
    if res["removed"]:
        L += ["", "## Duplicates removed", "", "| Source row | Reason | Kept row |", "|---|---|---|"]
        L += [f"| {r['source_row']} | {r['removed_reason']} | {r['kept_source_row']} |" for r in res["removed"][:25]]
        if n_dup > 25:
            L.append(f"\n_...and {n_dup - 25} more (see the Removed Duplicates sheet)._")
    if res["issues"]:
        L += ["", "## Needs a human look", "", "| Source row | Column | Value | Issue |", "|---|---|---|---|"]
        L += [f"| {i['source_row']} | {i['column']} | {i['value']} | {i['issue']} |" for i in res["issues"][:25]]
    L += ["", "## Assumptions", "",
          "- Dates read day-first unless `--monthfirst` was used.",
          "- Currency symbols were stripped from number columns.",
          "- Duplicate rules: " + "; ".join(["exact row match"] + [" + ".join(k) for k in res["key_specs"]]) + "."]
    out_path.write_text("\n".join(L) + "\n", encoding="utf-8")


def main(argv=None):
    p = argparse.ArgumentParser(description="Clean messy CSV/Excel files, remove duplicates, produce a report.")
    p.add_argument("input", type=Path)
    p.add_argument("--sheet", help="Excel sheet name or index (default: first)")
    p.add_argument("--keys", help='Duplicate keys, e.g. "email,first_name+last_name" (default: the email column)')
    p.add_argument("--phone-prefix", default="", help="Country prefix for local numbers, e.g. +92")
    p.add_argument("--monthfirst", dest="dayfirst", action="store_false", help="Read 03/04/2025 as March 4")
    p.add_argument("--outdir", type=Path, default=Path("output"))
    opts = p.parse_args(argv)

    if not opts.input.exists():
        sys.exit(f"File not found: {opts.input}")
    res = clean_file(opts.input, opts)
    opts.outdir.mkdir(parents=True, exist_ok=True)
    stem = opts.input.stem
    write_xlsx(res, opts.outdir / f"{stem}_clean.xlsx")
    res["clean"].drop(columns="source_row").to_csv(opts.outdir / f"{stem}_clean.csv", index=False, encoding="utf-8-sig")
    write_markdown(res, opts.outdir / f"{stem}_report.md")

    print(f"{opts.input.name}: {res['n_source']} rows -> {len(res['clean'])} clean "
          f"({len(res['removed'])} duplicates, {res['n_blank']} blank/footer, {len(res['issues'])} flagged)")
    print(f"Wrote {opts.outdir}/{stem}_clean.xlsx, _clean.csv, _report.md")


if __name__ == "__main__":
    main()
